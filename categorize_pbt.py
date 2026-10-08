"""
One-time PBT categorisation of listed companies from company_universe.xlsx.

For each Listed company: Screener se financials lo, calculate_csr_budget() wala
3-year average PBT nikalo (current FY exclude, uske pehle ke 3 FY), uska 2% CSR budget,
aur us CSR budget se category do:
    A: CSR budget > 1000 Cr | B: 500-1000 | C: 100-500 | D1-D5: 0-100 (pbt_category dekho)
    Loss Making: budget < 0 | NA: data nahi mila

Resumable: har company ka result company_pbt_cache.jsonl mein append hota hai,
dobara chalane par already-done companies skip ho jaati hain.

Usage:
    python categorize_pbt.py            # saari listed companies
    python categorize_pbt.py --limit 10 # sirf pehli 10 (testing)
    python categorize_pbt.py --retry-na # sirf NA wali companies dobara try karo
                                        # (NSE/BSE code lookup + standalone fallback)
    python categorize_pbt.py --partial-years  # NA wali (Screener page hai) pe available-years
                                              # rule: jitne recent saal ka PBT hai uska average
    python categorize_pbt.py --reclass  # cached CSR budget se class dobara nikalo (koi fetch nahi)
"""
import argparse
import csv
import difflib
import io
import json
import os
import re
import time

import openpyxl
import requests

from bs4 import BeautifulSoup

from financial_extractor import calculate_csr_budget
from search_tool import (SCREENER_HEADERS, _clean_base_company_name,
                         _normalize_screener_company_name,
                         get_financials_from_screener)

UNIVERSE_FILE = "company_universe.xlsx"
CACHE_FILE = "company_pbt_cache.jsonl"
OUTPUT_FILE = "company_pbt.xlsx"
DELAY_SECONDS = 1.5


def pbt_category(csr_budget):
    """Class CSR budget (avg PBT ka 2%, Cr) se. D ke 5 hisse: D1 25-100, D2 5-25,
    D3 1-5, D4 0.1-1, D5 0-0.1; negative budget (avg PBT loss) = "Loss Making"."""
    if csr_budget is None:
        return "NA"
    if csr_budget > 1000:
        return "A"
    if csr_budget >= 500:
        return "B"
    if csr_budget >= 100:
        return "C"
    if csr_budget < 0:
        return "Loss Making"
    if csr_budget >= 25:
        return "D1"
    if csr_budget >= 5:
        return "D2"
    if csr_budget >= 1:
        return "D3"
    if csr_budget >= 0.1:
        return "D4"
    return "D5"


def load_listed_companies():
    wb = openpyxl.load_workbook(UNIVERSE_FILE, read_only=True)
    ws = wb["Company Universe"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    i_name, i_isin, i_status = (header.index("Company Name"), header.index("ISIN"),
                                header.index("Listing Status"))
    companies = [{"company_name": r[i_name], "isin": r[i_isin] or ""}
                 for r in rows if r[i_name] and r[i_status] == "Listed"]
    wb.close()
    return companies


def load_cache():
    done = {}
    if os.path.exists(CACHE_FILE):
        with open(CACHE_FILE, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    rec = json.loads(line)
                    done[rec["company_name"]] = rec
    return done


def screener_search(query):
    """Screener search API only (Google/Serper fallback skip - credits bachane ke liye)."""
    for attempt in range(3):
        resp = requests.get("https://www.screener.in/api/company/search/",
                            params={"q": query}, headers=SCREENER_HEADERS, timeout=10)
        if resp.status_code == 429:
            print("[rate-limit] Screener 429, 60s wait...")
            time.sleep(60)
            continue
        resp.raise_for_status()
        results = resp.json()
        if results:
            return results[0]["name"], "https://www.screener.in" + results[0]["url"]
        return None, None
    raise RuntimeError("Screener rate limit - retries exhausted")


def find_on_screener(company):
    queries = [company["isin"], _normalize_screener_company_name(company["company_name"]),
               _clean_base_company_name(company["company_name"])]
    for q in dict.fromkeys(q for q in queries if q and len(q) >= 3):
        name, url = screener_search(q)
        if url:
            return name, url
    return None, None


def fetch_exchange_rows():
    """Official NSE + BSE equity lists -> [{"isin", "names", "code"}]. code = NSE
    symbol ya BSE scrip code; Screener pages /company/<code>/ pe khulte hain."""
    import extract_company as ec
    rows = []
    nse = ec.make_session(warm_up_url="https://www.nseindia.com")
    resp = nse.get("https://nsearchives.nseindia.com/content/equities/EQUITY_L.csv", timeout=30)
    for row in csv.DictReader(io.StringIO(resp.text)):
        row = {k.strip(): (v or "").strip() for k, v in row.items()}
        rows.append({"isin": row.get("ISIN NUMBER", ""), "code": row["SYMBOL"],
                     "names": [row.get("NAME OF COMPANY", "")]})
    bse = ec.make_session(warm_up_url="https://www.bseindia.com")
    bse.headers.update({"Referer": "https://www.bseindia.com/"})
    resp = bse.get("https://api.bseindia.com/BseIndiaAPI/api/ListofScripData/w"
                   "?Group=&Scripcode=&industry=&segment=Equity&status=Active", timeout=30)
    for row in resp.json():
        rows.append({"isin": (row.get("ISIN_NUMBER") or "").strip(), "code": str(row["SCRIP_CD"]),
                     "names": [(row.get("Issuer_Name") or "").strip(),
                               (row.get("Scrip_Name") or "").strip()]})
    return rows


def load_exchange_codes():
    """ISIN / lowercase name -> [NSE symbol, BSE scrip code]. Screener pages
    /company/<code>/ khulte hain, to search ke bina seedha page mil jaata hai."""
    codes = {}
    for row in fetch_exchange_rows():
        for key in [row["isin"]] + [n.lower() for n in row["names"]]:
            if key:
                codes.setdefault(key, []).append(row["code"])
    print(f"[codes] {len(codes)} NSE/BSE keys loaded")
    return codes


def _name_key(name):
    """Comparison ke liye: lowercase, & -> and, Ltd/Limited/(CN)/punctuation hatao."""
    name = re.sub(r"\(cn\)", " ", (name or "").lower()).replace("&", " and ")
    name = re.sub(r"\b(limited|ltd|the)\b", " ", name)
    return re.sub(r"[^a-z0-9]", "", name)


def screener_page_name(url):
    """Screener company page ka H1 naam, ya None agar page nahi hai."""
    resp = requests.get(url, headers=SCREENER_HEADERS, timeout=15)
    if resp.status_code == 429:
        print("[rate-limit] Screener 429, 60s wait...")
        time.sleep(60)
        resp = requests.get(url, headers=SCREENER_HEADERS, timeout=15)
    if not resp.ok or "/company/" not in resp.url:
        return None
    h1 = BeautifulSoup(resp.content, "html.parser").find("h1")
    return h1.get_text(strip=True) if h1 else None


def find_on_screener_retry(company, codes):
    """NA retry lookup: NSE/BSE code -> '&' wala naam -> 2-word query (sirf close name match)."""
    for code in codes.get(company["isin"]) or codes.get(company["company_name"].strip().lower()) or []:
        url = f"https://www.screener.in/company/{code}/consolidated/"
        name = screener_page_name(url)
        if name:
            return name, url
    full = _normalize_screener_company_name(company["company_name"])
    amp = re.sub(r"\s+and\s+", " & ", full, flags=re.I)
    if amp != full:
        name, url = screener_search(amp)
        if url:
            return name, url
    words = [w for w in re.findall(r"[A-Za-z0-9]+", full)
             if w.lower() not in ("and", "the", "limited", "ltd")]
    if len(words) >= 2:
        name, url = screener_search(" ".join(words[:2]))
        if url and difflib.SequenceMatcher(None, _name_key(name),
                                           _name_key(company["company_name"])).ratio() >= 0.85:
            return name, url
    return None, None


def _fy_only(fin):
    """'Mar 2013 10m' jaise odd-period columns hatao, sirf FYxx rakho (newest 4)."""
    years = sorted((fy for fy in fin.get("fiscal_years") or [] if re.fullmatch(r"FY\d{2}", fy)),
                   reverse=True)[:4]
    return dict(fin, fiscal_years=years)


def process(company, codes=None):
    rec = {"company_name": company["company_name"], "isin": company["isin"],
           "screener_name": None, "screener_url": None, "pbt": {},
           "calc_years": [], "average_pbt": None, "csr_budget_2pct": None,
           "category": "NA", "note": None}

    if codes is None:
        name, url = find_on_screener(company)
    else:
        name, url = find_on_screener_retry(company, codes)
    if not url:
        rec["note"] = "Screener pe nahi mila"
        return rec
    rec["screener_name"], rec["screener_url"] = name, url

    fields = fetch_pbt_fields(url, fallback=codes is not None)
    if fields is None:
        return None  # transient error - cache mat karo, next run retry karega
    rec.update(fields)
    return rec


def _recent_fy():
    """Sabse purana FY jo abhi 'recent' maana jaaye: latest completed FY se ek pehle
    (Oct 2026 mein latest FY26 -> recent = FY25 ya naya)."""
    t = time.localtime()
    latest = t.tm_year % 100 if t.tm_mon >= 4 else t.tm_year % 100 - 1
    return f"FY{latest - 1:02d}"


def partial_pbt_average(pbt):
    """Available-years rule (jab 3 poore pichhle saal na hon - nayi listing, demerger,
    year-end change): latest FY chhod ke uske pehle ke max 3 saal ka average; sirf
    1 saal ho to wahi. Latest data purana ho (< _recent_fy()) to (None, [])."""
    years = sorted((fy for fy, v in (pbt or {}).items()
                    if re.fullmatch(r"FY\d{2}", fy) and v is not None), reverse=True)
    if not years or years[0] < _recent_fy():
        return None, []
    calc = years[1:4] or years[:1]
    return round(sum(pbt[fy] for fy in calc) / len(calc), 2), calc


def partial_fields(pbt):
    """partial_pbt_average() ko record fields mein badlo, ya None agar rule lagu nahi."""
    avg, calc = partial_pbt_average(pbt)
    if avg is None:
        return None
    return dict(calc_years=calc, average_pbt=avg, csr_budget_2pct=round(avg * 0.02, 2),
                category=pbt_category(round(avg * 0.02, 2)), pbt_basis="partial",
                note=f"Partial: {len(calc)} saal ka PBT average ({', '.join(calc)})")


def fetch_pbt_fields(url, fallback=True):
    """Screener page se PBT/Class fields. fallback=True: odd periods hatao,
    consolidated mein data kam ho to standalone try karo (aur ulta), aur dono mein
    3 poore saal na hon to available-years rule. Transient error par None."""
    urls = [url]
    if fallback:
        urls.append(url.replace("consolidated/", "") if "/consolidated/" in url
                    else url.rstrip("/") + "/consolidated/")
    fields = None
    pages = []
    for i, page_url in enumerate(urls):
        fin = get_financials_from_screener(page_url, years=6 if fallback else 4)
        if fin.get("error"):
            return None
        if fallback:
            fin = _fy_only(fin)
        pages.append((page_url, fin.get("pbt") or {}))
        csr = calculate_csr_budget(fin)
        if i == 0 or csr["average_pbt"] is not None:
            fields = dict(screener_url=page_url, pbt=fin.get("pbt") or {},
                          calc_years=csr["calc_years"], average_pbt=csr["average_pbt"],
                          csr_budget_2pct=csr["csr_budget_2pct"],
                          category=pbt_category(csr["csr_budget_2pct"]), note=csr["note"])
        if csr["average_pbt"] is not None:
            return fields
        time.sleep(DELAY_SECONDS)
    if fallback:
        for page_url, pbt in pages:
            partial = partial_fields(pbt)
            if partial:
                return dict(partial, screener_url=page_url, pbt=pbt)
    return fields


def write_output(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PBT Categories"
    ws.append(["Company Name", "ISIN", "Screener Name", "Screener URL",
               "PBT Yr1 (Cr)", "PBT Yr2 (Cr)", "PBT Yr3 (Cr)", "Years Used",
               "Avg PBT (Cr)", "CSR Budget 2% (Cr)", "Class", "Note"])
    for rec in sorted(records, key=lambda r: r["company_name"].lower()):
        years = rec["calc_years"] or []
        pbts = [rec["pbt"].get(fy) for fy in years] + [None] * (3 - len(years))
        ws.append([rec["company_name"], rec["isin"], rec["screener_name"], rec["screener_url"],
                   *pbts[:3], ", ".join(years), rec["average_pbt"], rec["csr_budget_2pct"],
                   rec["category"], rec["note"]])
    wb.save(OUTPUT_FILE)

    counts = {}
    for rec in records:
        counts[rec["category"]] = counts.get(rec["category"], 0) + 1
    print(f"[excel] wrote {len(records)} rows -> {OUTPUT_FILE} | "
          + " ".join(f"{k}={v}" for k, v in sorted(counts.items())))


def reclass(rec):
    """Record ki class (aur verify ki class_before) uske csr_budget_2pct se dobara;
    NA / Review waise hi rehte hain."""
    budget = rec.get("csr_budget_2pct")
    if budget is None:
        return rec
    new = dict(rec)
    if rec["category"] not in ("NA", "Review"):
        new["category"] = pbt_category(budget)
    if rec.get("class_before") not in (None, "NA", "Review"):
        new["class_before"] = pbt_category(budget)
    return new


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--retry-na", action="store_true")
    parser.add_argument("--partial-years", action="store_true")
    parser.add_argument("--reclass", action="store_true")
    args = parser.parse_args()

    companies = load_listed_companies()
    done = load_cache()
    codes = None
    if args.reclass:
        # Badli hui class wale records cache mein append; load_cache() last wala rakhta hai.
        changed = 0
        with open(CACHE_FILE, "a", encoding="utf-8") as cache:
            for name, rec in list(done.items()):
                new = reclass(rec)
                if new != rec:
                    done[name] = new
                    cache.write(json.dumps(new, ensure_ascii=False) + "\n")
                    changed += 1
        print(f"[reclass] {changed} records updated")
        names = {c["company_name"] for c in companies}
        write_output([r for r in done.values() if r["company_name"] in names])
        return
    if args.partial_years:
        # NA records jinka Screener page hai: dono pages dobara padho, available-years rule ke saath.
        todo = [r for c in companies
                if (r := done.get(c["company_name"])) and r["category"] == "NA" and r.get("screener_url")]
        if args.limit:
            todo = todo[:args.limit]
        print(f"[start] {len(todo)} NA companies with a Screener page (partial-years rule)")
        with open(CACHE_FILE, "a", encoding="utf-8") as cache:
            for n, rec in enumerate(todo, 1):
                try:
                    fields = fetch_pbt_fields(rec["screener_url"], fallback=True)
                except Exception as e:
                    print(f"[error] {rec['company_name']}: {e}")
                    fields = None
                if fields:
                    rec = dict(rec, **fields)
                    done[rec["company_name"]] = rec
                    cache.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    cache.flush()
                    print(f"[{n}/{len(todo)}] {rec['company_name']} -> "
                          f"avg PBT {rec['average_pbt']} -> {rec['category']} | {rec['note']}")
                time.sleep(DELAY_SECONDS)
        names = {c["company_name"] for c in companies}
        write_output([r for r in done.values() if r["company_name"] in names])
        return
    if args.retry_na:
        # Cache mein naya record append hota hai; load_cache() last wala rakhta hai.
        todo = [c for c in companies if done.get(c["company_name"], {}).get("category") == "NA"]
        codes = load_exchange_codes()
    else:
        todo = [c for c in companies if c["company_name"] not in done]
    if args.limit:
        todo = todo[:args.limit]
    print(f"[start] {len(companies)} listed, {len(todo)} to fetch"
          + (" (NA retry)" if args.retry_na else ""))

    with open(CACHE_FILE, "a", encoding="utf-8") as cache:
        for n, company in enumerate(todo, 1):
            try:
                rec = process(company, codes)
            except Exception as e:
                print(f"[error] {company['company_name']}: {e}")
                rec = None
            if rec:
                done[rec["company_name"]] = rec
                cache.write(json.dumps(rec, ensure_ascii=False) + "\n")
                cache.flush()
                print(f"[{n}/{len(todo)}] {rec['company_name']} -> "
                      f"avg PBT {rec['average_pbt']} -> {rec['category']}")
            time.sleep(DELAY_SECONDS)

    names = {c["company_name"] for c in companies}
    write_output([r for r in done.values() if r["company_name"] in names])


if __name__ == "__main__":
    main()
